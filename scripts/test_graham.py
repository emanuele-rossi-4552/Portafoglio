"""Confronta il punteggio Graham di Python con quello calcolato da Excel (file ricalcolato con LibreOffice)."""
import json, os, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import aggiorna_dati as A

cfg = json.load(open(os.path.join(HERE, "..", "config.json"), encoding="utf-8"))
xl = sys.argv[1]
wb = openpyxl.load_workbook(xl, data_only=True)
ws = wb["Graham"]; gp = wb["Graham_Parametri"]
aaa = gp["C4"].value
def v(r, col):
    x = ws[f"{col}{r}"].value
    return None if x in (None, "") else x
bad = 0
for r in range(7, 40):
    if not ws[f"A{r}"].value: continue
    raw = {"price": v(r, "G"), "eps_ttm": v(r, "N"), "eps_avg": v(r, "O"), "bvps": v(r, "P"), "rev": v(r, "Q"),
           "cr": v(r, "S"), "ic": v(r, "U"), "roe": v(r, "V"), "fcf": v(r, "W"), "pay": v(r, "X"),
           "divy": v(r, "Z"), "posy": v(r, "AA"), "yn": v(r, "AB"), "cagr": v(r, "AC"), "sector": v(r, "L"),
           "nd_ebitda": v(r, "AZ"), "fwd_var": v(r, "BA"), "dd52": v(r, "BB"), "ma200": v(r, "BC")}
    ris = A.graham_calcola(raw, cfg["graham"], aaa)
    exp = dict(score=v(r, "D"), verdetto=v(r, "E"), cover=v(r, "F"), valore=v(r, "H"), mos=v(r, "I"), al=v(r, "J") or "")
    ok = (abs((ris["punteggio"] or 0) - (exp["score"] or 0)) < 1e-6 and ris["verdetto"] == exp["verdetto"]
          and abs((ris["copertura"] or 0) - (exp["cover"] or 0)) < 1e-9
          and (ris["valore"] == exp["valore"] or abs(ris["valore"] - exp["valore"]) < 1e-6)
          and (ris["margine"] == exp["mos"] or abs(ris["margine"] - exp["mos"]) < 1e-9)
          and ris["allerte"].strip() == str(exp["al"]).strip())
    bad += not ok
    sc = ris["punteggio"]
    print(("OK  " if ok else "DIFF"), ws[f"C{r}"].value, f"py {sc:.4f} / xl {exp['score']:.4f}" if sc is not None else "-", ris["verdetto"], "|", ris["allerte"], "|" if not ok else "", exp["al"] if not ok else "")
print("righe diverse:", bad)
sys.exit(1 if bad else 0)
